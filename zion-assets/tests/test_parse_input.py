"""Parser regressions using generated documents and simulated extraction engines."""
from pathlib import Path
import importlib.util
import io
from contextlib import redirect_stderr
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
spec = importlib.util.spec_from_file_location('parse_input', SCRIPTS/'parse_input.py')
parser = importlib.util.module_from_spec(spec)
spec.loader.exec_module(parser)


class PDF:
    def __init__(self, pages): self.pages = pages
    def __enter__(self): return self
    def __exit__(self, *args): return False
    def __iter__(self): return iter(self.pages)


class ParserRegression(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.base = Path(self.temporary.name).resolve()
        self.assertEqual(self.base.parent, Path(tempfile.gettempdir()).resolve())
        self.addCleanup(self.temporary.cleanup)
        self.pdf = self.base/'input.pdf'
        self.pdf.write_bytes(b'fixture')

    def test_docx_tables_nested_cells_and_body_order(self):
        try:
            import docx
        except ImportError:
            self.skipTest('python-docx is optional')
        doc = docx.Document()
        doc.add_paragraph('开头')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0,0).text = '林岚'
        table.cell(0,1).text = '医生，左手有伤'
        nested = table.cell(0,1).add_table(rows=1, cols=1)
        nested.cell(0,0).text = '嵌套备注'
        doc.add_paragraph('结尾')
        file = self.base/'characters.docx'
        doc.save(file)
        text = parser.parse_docx(str(file))
        self.assertIn('林岚\t医生，左手有伤', text)
        self.assertIn('嵌套备注', text)
        self.assertLess(text.index('开头'), text.index('林岚'))
        self.assertLess(text.index('嵌套备注'), text.index('结尾'))

    def test_docx_horizontal_merge_is_not_duplicated(self):
        try:
            import docx
        except ImportError:
            self.skipTest('python-docx is optional')
        doc = docx.Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0,0).merge(table.cell(0,1)).text = '合并单元格'
        file = self.base/'merged.docx'
        doc.save(file)
        self.assertEqual(parser.parse_docx(str(file)).count('合并单元格'), 1)

    def test_empty_middle_pdf_engine_continues_to_last_engine(self):
        final_reader = SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: '最终回退文本')])
        modules = {'fitz': SimpleNamespace(open=lambda path: PDF([])),
                   'pdfplumber': SimpleNamespace(open=lambda path: PDF([])),
                   'PyPDF2': SimpleNamespace(PdfReader=lambda file: final_reader)}
        with patch.dict(sys.modules, modules):
            self.assertEqual(parser.parse_pdf(str(self.pdf)), '最终回退文本')

    def test_pdf_exceptions_continue_to_next_engine(self):
        def fail(path): raise ValueError('fixture cannot parse')
        modules = {'fitz': SimpleNamespace(open=fail),
                   'pdfplumber': SimpleNamespace(open=lambda path: PDF([SimpleNamespace(extract_text=lambda: '中间引擎文本')])),
                   'PyPDF2': None}
        with patch.dict(sys.modules, modules):
            self.assertEqual(parser.parse_pdf(str(self.pdf)), '中间引擎文本')

    def test_first_successful_pdf_engine_stops_fallback(self):
        def forbidden(path): raise AssertionError('later extractor should not run')
        modules = {'fitz': SimpleNamespace(open=lambda path: PDF([SimpleNamespace(get_text=lambda: '首选文本')])),
                   'pdfplumber': SimpleNamespace(open=forbidden), 'PyPDF2': None}
        with patch.dict(sys.modules, modules):
            self.assertEqual(parser.parse_pdf(str(self.pdf)), '首选文本')

    def test_all_empty_pdf_engines_report_no_text_not_missing_dependencies(self):
        modules = {'fitz': SimpleNamespace(open=lambda path: PDF([])),
                   'pdfplumber': SimpleNamespace(open=lambda path: PDF([])),
                   'PyPDF2': SimpleNamespace(PdfReader=lambda file: PDF([]))}
        with patch.dict(sys.modules, modules):
            text = parser.parse_pdf(str(self.pdf))
        self.assertIn('OCR', text)
        self.assertNotIn('未安装', text)

    def test_missing_pdf_dependencies_are_distinguished(self):
        with patch.dict(sys.modules, {'fitz': None, 'pdfplumber': None, 'PyPDF2': None}):
            self.assertIn('未安装', parser.parse_pdf(str(self.pdf)))

    def test_invalid_frame_intervals_fail_before_writing_or_running_ffmpeg(self):
        for value in (0, -1, float('nan'), float('inf'), -float('inf'), True):
            with self.subTest(value=value), patch.object(parser.subprocess, 'run') as run:
                output = self.base/'frames'
                with self.assertRaises(ValueError):
                    parser.extract_video_frames('fixture.mp4', str(output), value)
                self.assertFalse(output.exists())
                run.assert_not_called()

    def test_valid_frame_interval_produces_advancing_timestamps(self):
        def ffmpeg(cmd, **kwargs):
            Path(cmd[-1]).write_bytes(b'frame')
            return SimpleNamespace(returncode=0)
        with patch.object(parser, 'get_video_metadata', return_value={'duration_seconds': 3.2}), patch.object(parser.subprocess, 'run', side_effect=ffmpeg):
            result = parser.extract_video_frames('fixture.mp4', str(self.base/'frames'), 1.5)
        self.assertEqual([item['timestamp'] for item in result], [0, 1.5, 3.0])
        self.assertEqual(len({item['path'] for item in result}), 3)

    def test_cli_rejects_invalid_frame_interval(self):
        video = self.base/'fixture.mp4'
        video.write_bytes(b'fixture')
        for value in ('0', '-1', 'nan', 'inf'):
            result = subprocess.run([sys.executable, '-B', '-X', 'utf8', str(SCRIPTS/'parse_input.py'), str(video), '--frame-interval', value, '--video-frames', str(self.base/'frames')], capture_output=True, timeout=10)
            self.assertEqual(result.returncode, 2)
        self.assertFalse((self.base/'frames').exists())

    def test_cli_output_cannot_replace_input(self):
        source = self.base/'scene.txt'
        source.write_text('林岚：签了。', encoding='utf-8')
        before = source.read_bytes()
        for output in (source, source.parent/'.'/source.name):
            result = subprocess.run([sys.executable, '-B', '-X', 'utf8', str(SCRIPTS/'parse_input.py'), str(source), '--output', str(output)], capture_output=True)
            self.assertEqual(result.returncode, 2)
            self.assertEqual(source.read_bytes(), before)

    def test_file_identity_alias_is_rejected_before_writing(self):
        source, alias = self.base/'scene.txt', self.base/'alias.txt'
        source.write_text('source body', encoding='utf-8')
        alias.write_text('alias fixture', encoding='utf-8')
        # Restricted Windows hosts may prohibit creating actual hardlinks.
        with patch.object(sys, 'argv', ['parse_input.py', str(source), '--output', str(alias)]), patch.object(parser.os.path, 'samefile', return_value=True), redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as result:
            parser.main()
        self.assertEqual(result.exception.code, 2)
        self.assertEqual(source.read_text(), 'source body')
        self.assertEqual(alias.read_text(), 'alias fixture')

    def test_cli_text_result_retains_original_source_name(self):
        import json
        source = self.base/'scene.txt'
        source.write_text('林岚：签了。', encoding='utf-8')
        result = subprocess.run([sys.executable, '-B', '-X', 'utf8', str(SCRIPTS/'parse_input.py'), str(source)], capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode, 0)
        data = json.loads(result.stdout)
        self.assertEqual(data['source_file'], 'scene.txt')
        self.assertEqual(data['text'], '林岚：签了。')


if __name__ == '__main__':
    unittest.main()
